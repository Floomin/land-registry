"""Формування Excel-файлів."""
import io

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .config import ALL, LEVEL_STYLES, SUMMARY_BUCKETS

THIN_SIDE = Side(style="thin", color="B7B7B7")
CELL_BORDER = Border(left=THIN_SIDE, right=THIN_SIDE, top=THIN_SIDE, bottom=THIN_SIDE)
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def to_excel(sheets: dict[str, pd.DataFrame]) -> bytes:
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="xlsxwriter") as w:
        for name, frame in sheets.items():
            frame.to_excel(w, sheet_name=name[:31], index=False)
            w.sheets[name[:31]].autofit()
    return buf.getvalue()


def make_slug(*parts: str) -> str:
    return "_".join(
        str(x).strip().replace(" ", "").replace("/", "-") for x in parts if x != ALL
    ) or "export"


def safe_sheet_name(name: str) -> str:
    for ch in '\\/*?:[]':
        name = name.replace(ch, "-")
    return (name.strip() or "Sheet")[:31]


def _write_summary_row(ws, r: int, label: str, values, level: int,
                       override_fill: str | None = None):
    style = LEVEL_STYLES[level]
    fill_color = override_fill or style["fill"]
    fill = PatternFill("solid", fgColor=fill_color) if fill_color else None
    font = Font(bold=style["bold"], color=style["font_color"], size=style["size"])

    lbl = ws.cell(row=r, column=1, value=label)
    lbl.font = font
    lbl.border = CELL_BORDER
    lbl.alignment = Alignment(indent=level, vertical="center")
    if fill:
        lbl.fill = fill

    for i, col_name in enumerate(SUMMARY_BUCKETS + ["Разом"], start=2):
        c = ws.cell(row=r, column=i, value=round(float(values[col_name]), 2))
        c.number_format = "#,##0.00"
        c.font = font
        c.border = CELL_BORDER
        c.alignment = Alignment(horizontal="right", vertical="center")
        if fill:
            c.fill = fill

    ws.row_dimensions[r].outlineLevel = level


def build_summary_workbook(pivot: pd.DataFrame) -> bytes:
    """Один аркуш на підприємство: Філія → Тип → Відділення → Сільська рада → Контрагент."""
    wb = Workbook()
    wb.remove(wb.active)
    headers = ["Ієрархія", *SUMMARY_BUCKETS, "Разом"]
    totals_cols = SUMMARY_BUCKETS + ["Разом"]

    for main_org, org_df in pivot.groupby("MainOrg", sort=True):
        ws = wb.create_sheet(title=safe_sheet_name(str(main_org)))
        ws.sheet_properties.outlinePr.summaryBelow = False
        ws.sheet_view.showGridLines = False

        header_fill = PatternFill("solid", fgColor="305496")
        for i, h in enumerate(headers, start=1):
            c = ws.cell(row=1, column=i, value=h)
            c.font = Font(bold=True, color="FFFFFF", size=11)
            c.fill = header_fill
            c.alignment = Alignment(horizontal="center", vertical="center")
            c.border = CELL_BORDER
        ws.row_dimensions[1].height = 22

        r = 2
        _write_summary_row(ws, r, f"{main_org} — усього", org_df[totals_cols].sum(),
                           level=0, override_fill="FFC000")
        r += 1

        for farm, farm_df in org_df.groupby("Farm", sort=True):
            _write_summary_row(ws, r, farm, farm_df[totals_cols].sum(), level=0)
            r += 1
            for typ, type_df in farm_df.groupby("Type", sort=True):
                _write_summary_row(ws, r, typ, type_df[totals_cols].sum(), level=1)
                r += 1
                for vid, vid_df in type_df.groupby("Відділення", sort=True):
                    _write_summary_row(ws, r, vid, vid_df[totals_cols].sum(), level=2)
                    r += 1
                    for council, council_df in vid_df.groupby("CouncilName", sort=True):
                        _write_summary_row(ws, r, council, council_df[totals_cols].sum(), level=3)
                        r += 1
                        for _, row in council_df.sort_values("Разом", ascending=False).iterrows():
                            _write_summary_row(ws, r, row["Контрагент"], row, level=4)
                            r += 1

        ws.freeze_panes = "B3"
        ws.column_dimensions["A"].width = 48
        for i in range(2, len(headers) + 1):
            ws.column_dimensions[get_column_letter(i)].width = 14

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
